#!/usr/bin/env python3
"""Build the editable product introduction from reviewed implementation evidence.

Run with the document workspace's bundled Python. Rendering and visual review
remain separate release checks.
"""
from pathlib import Path
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'
ASSETS=ROOT/'docs/evidence/screenshots'
OUT.mkdir(exist_ok=True);ASSETS.mkdir(parents=True,exist_ok=True)
for name in ['cloud-1440.png','editor-1440.png','screen-1920.png','edge-390.png']:
    source=ROOT/'.local/evidence'/name
    if source.exists() and not (ASSETS/name).exists():shutil.copy2(source,ASSETS/name)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)

def set_font(style,size,bold=False):
    style.font.name='Heiti SC';style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Heiti SC')
for name,size,bold in [('Normal',10.5,False),('Title',27,True),('Subtitle',13,False),('Heading 1',18,True),('Heading 2',13,True),('Caption',9,False)]:
    set_font(doc.styles[name],size,bold)
    doc.styles[name].paragraph_format.space_after=Pt(8)
    for border in doc.styles[name].element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    for fonts in doc.styles[name].element.findall('.//'+qn('w:rFonts')):
        for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attribute),None)
doc.styles['Normal'].paragraph_format.line_spacing=1.3
doc.styles['Normal'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(3)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(11)
doc.styles['Caption'].paragraph_format.line_spacing=1.15
header=sec.header.paragraphs[0];header.text='SmartFactory  产品介绍预稿';header.style=doc.styles['Caption']
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('2026 年 9 月 20 日    ');run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,level=2)
def page(title):
    doc.add_page_break();doc.add_heading(title,level=1)
def picture(name,width,caption):
    para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(ASSETS/name),width=Cm(width))
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(caption,'Caption')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for edge,size in [('top',100),('left',110),('bottom',100),('right',110)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(size));el.set(qn('w:type'),'dxa');margin.append(el)
            tcp.append(margin)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF1');tcp.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.18
                for run in para.runs:run.font.size=Pt(9.5);run.bold=ri==0
    p('')
    return t

p('SmartFactory 产品介绍','Title')
p('面向工厂监测与现场联动的云边平台','Subtitle')
p('本预稿面向赛事评审与学校材料审核，介绍产品用途、主要功能、实现基础和已有仿真结果。平台以 ThingsBoard CE、ThingsBoard Edge 和 DataTransfer 为基础，提供独立业务页面、控制会签、历史修订及 AI 草稿辅助。')
p('目前，五个赛事场景已通过实际协议模拟设备完成状态转换检查，500 台设备的采集性能测试持续运行了 60 分钟。完整验收仍需在统信 UOS Server V20 目标虚拟机中完成部署与长时间断云测试。')
picture('cloud-1440.png',16.5,'云端控制台实际截图。页面显示观测、来源状态、待投递数量与质量信息。')
h('使用对象与主要工作')
p('现场人员批准设备与采集配置，使用手持页面检查运行状态；工程师编写分析、告警与联动预案，管理人员按组织授权完成会签。大屏汇总指标和告警，数据异常及补传状态随观测一同展示。')

page('五个赛事场景')
p('示例工厂包含温湿度、红外照明、危气、AGV 和货物计数设备。模拟器实现实际协议服务，控制命令会改变设备后续读数，云端与边缘通过同一业务身份查看结果。')
table(['场景','判断与控制','可查看的结果'],[
('温湿度监控','阈值、回差与联锁判断后调节温湿度','读数恢复、执行回执和质量异常'),
('红外感应照明','有人点亮，无人延时关闭，抑制信号抖动','照明状态与命令去重结果'),
('危气监测','烟雾、可燃气和 CO 触发告警与排风预案','排风状态、跨节点进度与降级原因'),
('AGV 避障','距离不足时停车，恢复条件满足后继续','停车反馈、状态过期与异常处理'),
('货物感应计数','稳定事件身份去重，保存累计结果','重复、重启及断连后的准确计数')],[3.2,7.2,6.6])
h('一次现场操作的完整过程')
for text in [
'1. 设备通过 Modbus TCP、MQTT 或 OPC-UA 上报读数，DataTransfer 保存消息身份和协议质量。',
'2. 已发布定义计算条件；即时控制核对工程师、领导及适用的现场安全员确认。',
'3. 所属边缘节点检查当前数据与物理联锁，记录步骤并下发设备命令。',
'4. 模拟设备更新状态，平台保存回执，页面展示执行结果及相关数据快照。']:
    p(text)
h('现有场景证据')
p('五场景自动检查共完成 16 个状态转换。计数可靠性检查新增 10000 个有效事件，经过重复投递、网关重启和断连恢复后，云端与边缘的原始及派生累计值均为 20100。运行前已有的 10100 个事件同样保留。')
p('证据文件：docs/evidence/five-scenes.json；docs/evidence/counter-reliability.json。','Caption')

page('云端与现场的服务组织')
p('每个客户拥有独立实例和一个业务租户，云端与每个边缘节点分别保存业务数据库。现场静态页面、既有权限和已发布预案随边缘服务部署，断云时仍可由本地资源支撑操作。')
table(['位置','组成','职责'],[
('云端','ThingsBoard CE 与 Go 业务服务','实体与遥测投影，组织权限、发布、审批、修订、审计和 AI/MCP'),
('配置中心','独立 Go 服务与数据库','程序参数 Schema、版本、动态生效及加密凭据'),
('每个边缘节点','DataTransfer、Go 边缘服务、ThingsBoard Edge、PostgreSQL','本地采集、设备批准、权限校验、规则判断与命令执行'),
('同一现场','三个 NATS JetStream 实例','状态交换、预案版本核对、执行权与步骤进度')],[3,6,8])
h('数据保存与查询')
p('业务存储追加保存原始值、精确整数、质量原因、时间来源及设备和资产版本，再向 ThingsBoard 投递遥测投影。较早观测采用分块压缩归档，查询与补算统一读取热数据和归档，保留完整原始字段。')
p('默认原始观测及原始精度派生结果保存 30 天；分钟汇总保存 90 天，小时汇总保存 1 年，日汇总及告警历史保存 3 年。已发布版本和审计责任链长期保存，管理员可修改保留参数。')
h('现场协作与恢复')
p('跨节点执行通过三个协调副本取得多数确认，设备命令由所属节点下发。预案明确配置参与节点失效后的行为；失去多数或现场数据过期时，节点执行本地降级步骤。每个动作保留命令身份和步骤结果，恢复时核对既有执行记录。')
p('架构源码与接口说明见 docs/implementation/architecture.md、contracts/1.0 和 docs/implementation/development.md。','Caption')

page('三类编排与历史修订')
picture('editor-1440.png',17,'分析编辑器实际截图。示例将温度输入送入一分钟平均值窗口，并持久保存输出。')
p('分析、告警与规则策略共用版本化定义及编辑器，各类定义提供适用节点。工程师能够保存草稿、校验类型和依赖、使用样本模拟，再从评审页面发布；既有正式版本与草稿分别保存。')
h('质量与计算依据')
p('默认只有 GOOD 数据参与业务计算，BAD 与 UNCERTAIN 保存原始记录及原因。页面显示排除数量、已知缺口、来源状态和完整程度；系统无法确定预期数量时显示未知。业务阈值超限由告警定义判断。')
h('补传与修订记录')
p('补传自动创建原始资料可用范围内的补算任务，使用对应历史时段的规则及资产版本。补算执行分析与告警逻辑，保存旧值、新值和原因；图表提供修订标记与对比，失败任务保留影响范围。控制动作由实时路径及边缘调度执行。')
h('AI 与外部 MCP')
p('页面助手与 MCP 按资源权限发现目录、读取结构、查询数据，并协助创建、修改、校验和比较三类草稿。服务端将 AI 身份限定为查询及草稿操作，正式发布和设备控制由有权人员操作。模型端点、名称、密钥与超时均可配置，固定模型模拟器提供可复现的流式对话与工具调用检查。')

page('现场页面与管理大屏')
picture('screen-1920.png',17,'管理大屏实际截图。指标、趋势、来源与补传状态均来自业务 API，可配置展示内容。')
p('云端控制台覆盖设备与资产、数据监控、三个编排入口、审批执行、组织授权、配置、插件、审计和运行状态。管理大屏提供指标详情与异常状态，现场页面使用移动布局呈现本地监测、设备批准和既有预案操作。')
table(['页面','检查分辨率','已执行检查'],[
('云端控制台','1440 × 900','布局、编排保存与评审流程、助手错误呈现'),
('现场页面','390 × 844','移动布局、设备批准与资源权限'),
('管理大屏','1920 × 1080','指标配置、精确大整数与详情查看')],[4,4,9])
p('浏览器功能记录共八项通过；20 个独立浏览器会话的显示延迟 p95 为 798 毫秒。目标环境将继续检查持续运行、断网、补传及来源异常时的页面状态。','Caption')

page('控制审批与责任追溯')
p('云端即时控制引用已发布预案，边缘在执行前核对当前条件。拒绝回执包含原因、风险及数据依据；预案参数、版本或风险依据变化后，需要重新完成适用确认。')
table(['操作','默认确认要求'],[
('普通即时控制','一名工程师和一名领导'),
('业务强制执行','两名工程师和一名领导'),
('安全强制执行','上述会签，加指定现场安全员的本地二次认证'),
('物理安全联锁','所有软件控制入口均检查联锁状态')],[5,12])
p('确认按不同人员身份计数，默认有效期 5 分钟。正式下发后允许开始执行的期限为 10 秒，步骤遵循各自超时设置；非幂等动作的发送结果无法确认时，执行进入结果待核对。')
h('离线授权与连接准入')
p('云边业务通信采用 mTLS，并校验证书对应节点的登记状态。签名离线权限包默认有效 48 小时，保存既有人员权限；新策略与权限提升在云端完成。组织变更、人员停用和角色撤销参与授权检查。')
h('审计与恢复')
p('审计保存操作时的身份、来源、时间、关联请求、版本和数据快照，按来源维护序号、哈希链及签名检查点。控制追溯关联审批、判断、命令和回执；备份恢复检查包含对象数量、版本、审计完整性及已执行命令的去重结果。')
h('开源基础与自研内容')
p('原生平台固定为 ThingsBoard CE v4.3.1.5 与 ThingsBoard Edge 4.3.1.1EDGE。项目自研内容包含 DataTransfer 扩展、Go 业务服务、独立配置中心、三套前端、公共契约、协议模拟器和验收工具；源码包保留上游许可证、依赖清单和修改说明。')

page('阶段测量与交付安排')
p('以下数据来自 2026 年 9 月 20 日开发环境的实测报告。开发主机为 macOS arm64，PostgreSQL、NATS 和原生平台采用锁定镜像的 Linux 容器；报告同时记录各次运行的二进制摘要与测试参数。')
table(['项目','实际结果'],[
('500 设备持续采集','60 分钟，18654020 条单测点消息；5181.65 条/秒，输入与保存数量一致'),
('资源与命令延迟','最大 RSS 58.5 MiB；1000 次控制测量，p99 为 18.584 ms'),
('持久消息故障注入','100000 个事件，丢确认 14285 次，进程重启 3 次；已确认缺失及重复业务效果均为 0'),
('API 并发测量','20 个 SSE 订阅、10 个历史查询客户端，持续 180 秒；历史 p95 1564.28 ms，SSE p95 1096 ms'),
('页面实际显示','20 个独立浏览器会话，435 个样本，显示延迟 p95 为 798 ms'),
('现场故障切换','三节点、150 台设备、100 次故障，持续 1333.94 秒，重复设备动作 0'),
('完整观测压缩归档','8192 点 PostgreSQL 样本，归档含索引及 TOAST 共 499712 字节，平均 61 字节/点')],[4.5,12.5])
p('采集 60 分钟报告对应当次二进制；后续批处理与归档修改另有回归记录。页面测量覆盖源观测时间至指标 DOM 更新后的两次动画帧。归档样本用于结构比较，实际 24 小时容量报告还需覆盖投影、修订、待投递数据与备份。','Caption')
h('目标环境待验工作')
p('统信 UOS Server V20 x86_64 虚拟机使用 8 vCPU、16GB 内存和 220GB 磁盘作为初始配置。待验项目包括完整安装与恢复、单节点和三节点连续断云 24 小时、100Mbps 与 20ms 链路下的补传补算，以及实际 24 小时存储测量。完整原始数据的容量预算将按上述实测调整。')
h('材料日期')
p('10 月 7 日前提交产品介绍和 PPT 预稿，10 月 18 日作为完整验收目标，10 月 19 日整理源码、部署包、文档与演示视频。最终提交暂按 10 月 20 日前准备，并随学校或发榜单位的新通知调整。')
p('量化证据分别见 docs/evidence/capacity-disk-60m.json、A06-durable-100000.json、api-postgres-180s.json、browser-latency-20.json、site-150-faults-100.json、archive-postgres-v2.json。完整追踪表见 docs/implementation/requirements.md。','Caption')

doc.core_properties.title='SmartFactory 产品介绍'
doc.core_properties.subject='赛事产品介绍预稿与阶段验证结果'
doc.core_properties.author='SmartFactory 项目'
doc.core_properties.keywords='云边协同,工业采集,控制审批,ThingsBoard'
target=OUT/'SmartFactory产品介绍预稿.docx'
doc.save(target)
print(target)
