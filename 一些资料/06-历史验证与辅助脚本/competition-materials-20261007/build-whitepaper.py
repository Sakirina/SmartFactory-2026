from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path('/path/to/SmartFactory-2026')
BUILD=ROOT/'.local/competition-materials-20261007'
ASSETS=ROOT/'output/playwright/competition-materials'
OUT=ROOT/'deliverables/SmartFactory产品白皮书.docx'
FONT='Heiti SC'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.3)
sec.different_first_page_header_footer=True

def setfont(run,size=None,bold=None,color='000000'):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        run._element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    return run

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name]
    style.font.name=FONT; style.font.color.rgb=RGBColor(0,0,0)
    style.font.italic=False
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        style._element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    style.paragraph_format.widow_control=True
normal=doc.styles['Normal']
normal.font.bold=False
doc.styles['Subtitle'].font.bold=False
doc.styles['Caption'].font.bold=False
normal.font.size=Pt(11.1)
normal.paragraph_format.line_spacing=1.32
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.space_before=Pt(0)
for name,size in [('Heading 1',20),('Heading 2',13.2),('Heading 3',11.6)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
    s.paragraph_format.space_before=Pt(12 if name!='Heading 1' else 0)
    s.paragraph_format.space_after=Pt(7 if name!='Heading 1' else 13)
    s.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_before=Pt(4)
doc.styles['Caption'].paragraph_format.space_after=Pt(9)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
setfont(footer.add_run('SmartFactory 产品白皮书  '),8.5,color='666666')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='SmartFactory 产品白皮书'
doc.core_properties.subject='工业现场数据与业务协作平台'
doc.core_properties.author='SmartFactory'
doc.core_properties.last_modified_by='SmartFactory'
doc.core_properties.keywords='SmartFactory,工业物联网,云边协同,设备采集,规则编排'
doc.core_properties.comments=''
doc.core_properties.created=datetime(2026,10,7,tzinfo=timezone.utc)
doc.core_properties.modified=datetime(2026,10,7,tzinfo=timezone.utc)

def p(text,style=None):
    pp=doc.add_paragraph(style=style)
    setfont(pp.add_run(text))
    return pp
def h(text,level=1): return doc.add_heading(text,level)
def page(title):
    pp=h(title);pp.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    values=[headers,*rows]
    for ri,vals in enumerate(values):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trPr=t.rows[ri]._tr.get_or_add_trPr()
        keep=OxmlElement('w:cantSplit'); trPr.append(keep)
        if ri==0:
            header=OxmlElement('w:tblHeader'); trPr.append(header)
        for ci,value in enumerate(vals):
            c=cells[ci]; c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            cp.append(borders)
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'183C4C' if ri==0 else ('F3F6F7' if ri%2==0 else 'FFFFFF')); cp.append(shd)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top',90),('bottom',90),('left',115),('right',115)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            q=c.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.16
            setfont(q.add_run(value),10.3,ri==0,'FFFFFF' if ri==0 else '000000')
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1
    setfont(spacer.add_run(''),1)
    return t
def pic(name,caption,width=6.94,maxheight=None):
    source=ASSETS/name
    with Image.open(source) as im: iw,ih=im.size
    if maxheight: width=min(width,maxheight*iw/ih)
    q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after=Pt(0);q.paragraph_format.keep_with_next=True
    shape=q.add_run().add_picture(str(source),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
def link(label,url):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(6)
    rel=pp.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    fs=OxmlElement('w:rFonts');fs.set(qn('w:ascii'),FONT);fs.set(qn('w:hAnsi'),FONT);fs.set(qn('w:eastAsia'),FONT);pr.append(fs)
    color=OxmlElement('w:color');color.set(qn('w:val'),'176B66');pr.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz)
    r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);hl.append(r);pp._p.append(hl)

# Cover
q=doc.add_paragraph(style='Title');q.paragraph_format.space_before=Pt(35);q.paragraph_format.space_after=Pt(15)
setfont(q.add_run('SmartFactory'),42,True)
q=doc.add_paragraph(style='Subtitle');q.paragraph_format.space_after=Pt(12)
setfont(q.add_run('产品白皮书'),27,True)
q=p('工业现场数据与业务协作平台');q.paragraph_format.space_after=Pt(22)
setfont(q.runs[0],16)
p('SmartFactory 将设备采集、现场规则和人员协作组织在同一套业务流程中。使用者可以沿着设备观测追踪分析结果、告警处置、控制审批、设备反馈与工单交接，并在规则变更、通信中断或程序更新后继续核对原有记录。')
pic('cloud-live.png','云端运行概览  五场景示例数据',6.94,4.2)
q=p('产品版本 1.1.0    2026 年 10 月');q.paragraph_format.space_before=Pt(12)
setfont(q.runs[0],10.5,color='4B4B4B')

page('产品定位与应用场景')
p('工业现场需要同时处理不同协议的设备、持续变化的测点数据和跨岗位的处置工作。SmartFactory 面向这类工作提供统一设备标识、数据质量记录、规则版本和业务状态，使工程师、现场操作人员与管理人员能够围绕同一设备、同一次异常和同一份执行记录开展工作。')
p('产品由云端控制台、现场工作台和管理大屏组成。云端维护全厂视图、人员与发布，现场服务就近采集和执行已发布规则，大屏按权限展示指标、趋势与来源状态；设备通过所属边缘节点进入平台，现场工作持续保留操作与反馈依据。')
h('五类工厂场景',2)
table(['应用场景','观测与判断','业务结果'],[
 ['温湿度与通风','温度、湿度、上下限与回差','环境分析、异常告警，以及通风和温湿度调节'],
 ['红外照明','人员存在、持续时间与延时','有人点亮，无人持续达到条件后关闭'],
 ['危气监测','烟雾、可燃气体与一氧化碳','分别判断告警条件并执行排风预案'],
 ['AGV 避障','距离、数据有效性与物理联锁','障碍停车，满足恢复条件后继续运行'],
 ['货物计数','脉冲事件和稳定事件身份','去除重复事件，保存有效累计计数'],
],[1.23,2.48,3.23])
h('使用者与工作方式',2)
p('工程师维护设备参数和规则，利用连续模拟及历史比较检查变更；现场人员核对设备、审批条件与实际反馈，并记录处理过程；管理人员通过全厂视图查看异常、责任归属和后续工作。权限随人员角色、组织关系、资源范围及当前业务状态共同计算。')
p('场景模板预先组织协议、测点、动作和规则关系。使用者绑定现场设备及连接器后生成草稿，在评审中确认参数和影响，随后采用新的规则版本；同类设备可以沿用模板结构，再按工艺条件调整具体参数。')

page('系统架构与技术组成')
p('SmartFactory 采用云端与边缘协同的部署结构。DataTransfer 负责协议设备，边缘业务服务保存现场观测、计算与控制，云端维护全厂业务和人员，配置中心按独立程序身份提供参数与密钥引用。每个部署实例使用自己的持久数据和运行身份。')
table(['组件','承担的工作'],[
 ['DataTransfer','Modbus TCP、MQTT、OPC UA 采集，统一观测、缓冲确认及命令结果'],
 ['sf-edge','设备准入、现场规则、控制状态、离线授权与云边同步'],
 ['sf-cloud','全厂查询、定义评审、业务处置、人员管理及发布协调'],
 ['sf-config 与发布代理','身份校验、固定配置分发、工件核验、进程切换与运行报告'],
 ['ThingsBoard CE 与 Edge','原生实体、遥测、规则链和告警投影'],
 ['PostgreSQL 与 Kafka','业务及原生数据库持久化，以及原生平台消息服务'],
 ['NATS JetStream 与 Envoy','多节点执行权和签名检查点，以及浏览器 HTTPS 转发'],
 ['独立前端','云端、现场、大屏页面和类型化业务 API 客户端'],
],[2.06,4.88])
h('观测与控制的处理过程',2)
p('设备观测进入 DataTransfer 后形成统一字段，边缘业务桥检查设备批准状态、节点归属及配置版本，再保存观测和投递记录。计算服务采用发布时编译的计划生成分析结果、告警或预案意图，现场控制继续检查人员许可、条件和设备反馈。')
p('云边通信通过双向 TLS 认证节点，使用持久消息身份、确认和游标交换数据与业务资料。业务更新将对象、审计和后续投递事项放在对应事务中保存，网络发送和原生平台投影在提交后继续处理。')
h('项目实现与开源组件',2)
p('项目在 ThingsBoard CE／Edge 的原生能力上实现独立的工业业务服务、DataTransfer 采集程序和三套页面。风险会签、历史规则与资产版本、控制恢复、工单交接、AI 调查及固定发布由 SmartFactory 业务层维护，组件来源和许可随部署包提供。')

page('设备采集与数据治理')
p('设备接入以统一设备标识和所属节点为依据。现场人员核对协议地址、采样周期、测点与动作映射并批准设备，业务桥随后保存有效观测；参数编辑采用协议元数据生成表单，并保留扩展字段、已阅版本和实际采用回执。')
table(['协议','采集与设备操作'],[
 ['Modbus TCP','寄存器与位字段映射、字节顺序、数据转换及反向编码'],
 ['MQTT','主题与载荷映射、状态或事件采集，按动作映射发送命令'],
 ['OPC UA','节点读取、订阅、方法调用，以及签名加密和证书配置'],
 ['进程插件','通过独立进程扩展协议和外部系统，保存调用及失败状态'],
],[1.35,5.59])
h('数值精度与质量语义',2)
p('原始观测同时保存值、单位、质量、时间和来源。大整数及高精度数值在数据表、工具提示与导出中保留原值，图表坐标采用显示近似值；GOOD、BAD 和 UNCERTAIN 描述观测质量，温度超过工艺阈值时仍可具有有效质量，异常判断由告警规则承担。')
p('查询结果一并展示已记录缺口、来源通信状态和补传进度，使用者可以据此判断某一时间范围的数据情况。服务端负责筛选与分页，完整导出沿同一查询快照读取后续页面，可恢复订阅继续追踪已提交的变化。')
h('保留策略与历史修订',2)
table(['数据层级','默认保留时间'],[
 ['原始观测','30 天，其中最近 24 小时保留在热表，较早内容分块归档'],
 ['分钟与小时汇总','分钟汇总 90 天，小时汇总 365 天'],
 ['日汇总与告警','1095 天'],
],[2.08,4.86])
p('热表和归档共用查询入口，保留期可通过配置调整。延迟观测在自动处理范围内按当时的规则和资产版本计算，修订记录保存前后值、原因和任务；历史补算沿分析与告警路径处理，设备动作由实时策略及控制流程执行。')

page('规则编排与历史分析')
p('数据分析、告警和联动预案共用定义结构，分别描述计算输出、异常判断和设备步骤。编辑器从节点目录读取参数类型、范围及单位，保存草稿后检查端口、依赖、表达式和循环关系；单份定义支持 1 至 128 个节点及最多 512 条连接。')
pic('editor-focus.png','可视化规则编辑器  一分钟平均温度示例',6.94,3.8)
h('连续模拟与语义评审',2)
p('连续模拟使用正式计算内核，固定草稿修订、输入观测、历史、资产版本和逻辑时钟，逐条显示节点输入输出、状态变化和动作意图。使用者可导出复现文件，再通过语义评审检查节点、参数、控制政策及资源引用的变化。')
h('回放比较与后台恢复',2)
p('历史回放采用已发布规则，两版本比较共享同一份输入和初始状态，影子候选将候选结果与对应正式结果配对。计算任务保存原业务身份、阶段和已提交游标，失败重试继续已有检查点，取消操作处理后续计算，已经提交的结果和解释继续保留。')
p('规则发布保存新的生效版本和编译计划。已有执行实例继续使用创建时固定的预案，恢复历史规则内容则生成新的草稿，再完成模拟和评审，使每次采用都有明确的来源与记录。')

page('控制审批与执行恢复')
p('人工控制请求固定预案、参数和风险依据，符合角色与组织关系的人员完成会签后正式下发。系统在下发和各设备步骤前核对授权、审批有效期、现场条件及物理联锁，并将命令、反馈和人工处理保存在同一执行记录中。')
table(['控制类别','默认审批条件'],[
 ['普通即时控制','一名工程师及一名符合发起人组织关系的领导'],
 ['业务强制执行','两名工程师及一名符合组织关系的领导'],
 ['安全强制执行','完成强制会签，再由指定安全员在目标现场通过二次认证确认'],
],[1.60,5.34])
p('每人在同一请求承担一个会签角色，默认会签及二次认证窗口均为五分钟，正式下发后的启动期限为十秒，具体数值由配置中心管理。参数、风险依据或人员许可发生变化时，系统按最新条件重新检查请求。')
h('结果未知时的核对过程',2)
for text in [
 '1. 打开执行详情，检查原命令、预案版本、设备步骤和当前允许操作，阅读已经保存的反馈来源。',
 '2. 提交核对原命令结果的请求，记录操作理由，沿原请求身份追踪受理、开始与终止回执。',
 '3. 结合观测时间、采集时间和设备状态确定已经完成的动作，满足条件后恢复剩余步骤。',
 '4. 取消执行时停止后续步骤，已经发送的命令与迟到反馈继续关联到原执行记录。',
]:p(text)
h('重复请求与多节点协作',2)
p('控制请求使用稳定的下发身份，各设备步骤具有固定命令身份，DataTransfer 在发送前保存动作日志。同一身份和相同内容的重复请求取得已有状态，结果未知时继续核对原命令；相同身份携带不同内容时返回冲突，交由人员核对。')
p('多节点预案通过 NATS 三副本保存执行权、代次和签名检查点。节点在规则、参与者或通信条件发生变化时，按预案配置进入本地降级步骤；恢复后结合检查点及实际设备效果继续处理。现场 PLC 与物理联锁仍按设备和工艺要求运行。')

page('告警处置与工单交接')
p('告警记录设备异常的发生和恢复周期，人工处置记录确认、责任人与处理状态。持续采样和人员操作分别维护业务版本，新的观测保持相同动作含义时，人员可以继续使用已阅依据；恢复、级别变化或人工处置改变条件后，系统要求重新核对。')
pic('work-detail-focus.png','工单详情保存责任人  关联告警和处理记录',6.94,3.65)
h('责任转移与已阅依据',2)
p('工单关联告警或控制执行，记录处理事项、责任人和进度。交接时，当前责任人或管理员指定新的责任人、说明未完成事项，并选择已经阅读的告警或执行依据；责任更新、交接请求、依据版本和审计一同保存。')
p('新的责任人根据当前工单状态取得允许操作，前一责任人的操作范围随交接结果更新。发生版本冲突时，页面保留填写内容，使用者阅读最新状态后继续提交，避免在旧资料上完成责任转移。')
h('场景模板与批量准备',2)
p('五类模板把协议、测点、动作和规则组织成可复用配置，一次准备支持 1 至 20 个场景实例。服务检查设备归属、连接器版本及操作许可，全部通过后在同一批次保存草稿与审计，后续按实例记录配置演进和发布采用情况。')

page('AI 调查与外部系统协作')
p('页面助手通过当前人员身份访问设备、数据、定义、执行与历史结果，帮助使用者定位异常的相关资料，并在具备草稿许可时编辑规则。外部 MCP 客户端调用同一应用服务，工具请求与页面操作采用一致的资源和业务许可。')
table(['能力','提供给使用者的内容'],[
 ['数据与对象查询','发现目录、读取字段结构，按设备和时间范围查询数据及业务记录'],
 ['实际调用依据','保存进入模型的工具 JSON、资源版本、内容摘要及调用状态'],
 ['草稿协作','创建或修改分析、告警、联动草稿，执行校验、模拟及差异比较'],
 ['持久调查','保存问题、答案、调用过程和终止状态，重新打开后继续查看'],
 ['模型配置','支持 Chat Completions 与 Responses，按供应商 API 配置端点和模型'],
],[1.45,5.49])
h('可追溯的调查过程',2)
for text in [
 '1. 使用者明确设备、测点、时间范围或规则问题，平台按当前身份建立独立调查。',
 '2. 模型调用工具取得业务资料，系统记录实际发送的原文、来源版本及预算使用情况。',
 '3. 使用者打开依据详情，核对引用与原对象，并结合本次调用结果阅读回答。',
 '4. 涉及规则调整时保存草稿，由具备发布许可的人员完成模拟、评审与采用。',
]:p(text)
h('权限与内容生命周期',2)
p('引用在读取时重新检查当前授权。依据默认保留七天，过期、删除或失去资源许可后，服务清除相应原文、参数和受影响答案；单个工具结果预算为 512 KiB，敏感凭据在发送与保存前处理，调查保留审计所需的身份、时间和内容摘要。')
h('外部集成',2)
p('公开 HTTP API 使用 /api/sf/v1，并提供 OpenAPI 3.1 和 JSON Schema。MCP 使用专用 AI 身份提供读取与草稿工具，组织数据可通过独立进程插件导入；现场设备协议、外部业务系统与 AI 客户端分别沿相应契约扩展。')

page('节点配置与部署运维')
p('节点发布将程序、规则、参数和连接器组成不可变清单，记录目标平台、内容摘要及数据库兼容条件。配置中心按工作负载身份提供固定内容，发布代理完成准备与核验后管理实际 Edge 子进程，云端根据运行报告推进后续批次。')
table(['发布对象','固定与核对内容'],[
 ['程序工件','应用版本、目标平台、文件摘要及兼容声明'],
 ['规则与配置','来源版本、内容摘要、节点本地采用版本与密钥引用'],
 ['节点身份','所属节点、工作负载类型、凭据代次及授权状态'],
 ['批次进度','当前批次、目标节点、准备结果、运行报告及失败原因'],
],[1.45,5.49])
h('分批运行与失败恢复',2)
p('代理切换进程时保留原数据库和主密钥。协调器等待前一批次的有效运行报告，再允许下一批次继续；失败后的恢复沿原任务和经过授权的新代次进行，页面保留处理理由、来源版本和当前允许操作。程序回退涉及数据库兼容范围时，按配套备份恢复流程处理。')
h('部署与数据保护',2)
p('持久部署采用 Linux x86_64、Docker Compose 和 PostgreSQL，提供单节点与三节点部署配置，并记录自托管或托管运维归属。三节点现场配置使用三个 NATS 副本，资源规划同时计入原生平台、业务数据库、消息服务、归档与备份。')
p('安装工具为实例生成独立私有目录、随机凭据、客户 CA 和节点身份。完整备份保存业务及原生数据库、消息卷、采集状态、配置和密钥；保留原目录与数据卷的再次启动和全量恢复分别按照对应步骤执行。')
h('目标部署环境',2)
p('当前发布包面向 Linux amd64。UOS Server V20 x86_64 是目标系统，完整安装恢复、单节点和三节点连续断云 24 小时、受限链路补传补算，以及实际 24 小时容量测量仍待在该目标环境执行，实施时据此确定容量和运行参数。')

page('使用界面与现场演示')
p('三套页面采用同一业务 API 和资源许可。云端控制台提供全厂查询与管理，现场工作台读取所属节点的数据和有效离线授权，大屏按配置展示指标、趋势和来源状态；页面中的数值可以继续追踪到数据记录及对应业务对象。')
pic('screen-live.png','管理大屏  指标  趋势与现场来源状态',6.94,3.65)
table(['页面','主要工作'],[
 ['云端控制台','设备与资产、分析编排、告警工单、历史调查和分批发布'],
 ['现场工作台','设备批准、现场观测、控制确认、反馈核对及本地业务处理'],
 ['管理大屏','集中查看指标、趋势、告警及来源通信与补传状态'],
],[1.45,5.49])
h('温湿度场景演示流程',2)
for text in [
 '1. 在现场确认温湿度设备及连接器，查看测点数值、采样时间与质量。',
 '2. 打开温度分析草稿，运行连续样本模拟，检查阈值、回差和节点解释，再完成规则评审。',
 '3. 观察模拟设备触发告警，查看告警详情、责任人及关联工单，记录现场处理或交接事项。',
 '4. 按预案完成控制审批，核对原命令与设备反馈，再从历史数据检查处理前后的变化。',
]:p(text)
p('本白皮书的页面配图来自应用实际界面，展示本机五场景示例或独立业务示例数据。')

page('验证结果与测试条件')
p('1.1 系列围绕规则、存储、持久任务、控制、业务协作、AI 与节点发布完成五阶段共 45 项功能验收。浏览器、数据库、协议及独立程序组合分别记录受测输入与运行条件，涉及共享调用的变更继续核对相关结果。')
table(['验证范围','已完成的结果'],[
 ['功能验收','45 项，覆盖五阶段业务能力'],
 ['独立浏览器检查','39 项，覆盖参数、规则、告警工单、查询、AI 和节点发布等流程'],
 ['Linux 包入口检查','55 项，包含独立 PostgreSQL、代理节点、实际 DataTransfer、缓存启动及原数据再次启动'],
 ['协议与原生专项','OPC UA 签名加密、MQTT、进程插件、云边 mTLS、消息故障及原生组件配套恢复'],
],[1.60,5.34])
h('设备采集性能基线',2)
p('2026 年 9 月 20 日的独立 DataTransfer 测量在 macOS arm64、10 个逻辑 CPU 环境执行，采用 200 台 Modbus、200 台 MQTT 和 100 台 OPC UA 模拟设备，预热 60 秒后持续测量 60 分钟，每条消息包含一个测点。该记录对应当次采集程序构建，供后续目标环境测试比较。')
table(['测量指标','记录结果'],[
 ['累计输入与保存数量','均为 18,654,020 条单测点消息'],
 ['平均处理速率','5,181.65 条单测点消息／秒'],
 ['采集进程最大 RSS','58.50 MiB'],
 ['控制延迟','1,000 次测量的 p99 为 18.584 ms'],
],[2.40,4.54])
h('环境与持续时长',2)
p('原生 CE／Edge 组合在 aarch64 Docker 主机上运行 linux/amd64 镜像，相关记录包含架构模拟、峰值资源和停止恢复情况。UOS 全系统安装恢复、24 小时断云及容量测量继续按目标环境专项执行；现场容量规划以部署组件、设备频率、保留时间和专项实测共同确定。')

page('交付内容与技术资料')
p('SmartFactory 1.1.0 提供 Linux amd64 程序归档及 SHA-256 校验文件，包内包括独立程序、三套前端、源码、接口契约、部署脚本、使用文档和组件许可。运行镜像按固定摘要另行准备，部署工具根据所选形态生成实例配置和身份资料。')
table(['技术组成','1.1 系列固定版本'],[
 ['ThingsBoard CE 与 Edge','4.3.1.6 ／ 4.3.1.6EDGE'],
 ['数据库与消息','PostgreSQL 18.6、Kafka 4.3.1、NATS 2.14.7'],
 ['服务转发','Envoy 1.39.1'],
 ['前端','React 19.2.5、Vite 8.1.5、ECharts 6.1.0'],
 ['构建工具','Go 1.27.1、Node.js 24.21.0'],
],[2.1,4.84])
h('资料与版本来源',2)
p('产品使用、技术实现与验证范围可以从下列资料继续查询。源码中的部署清单、API 契约和组件来源记录与版本包配套保存，开源组件保留原有来源和许可说明。')
base='https://github.com/Sakirina/SmartFactory-2026/blob/v1.1.0/'
for label,url in [
 ('SmartFactory 1.1.0 程序与校验文件','https://github.com/Sakirina/SmartFactory-2026/releases/tag/v1.1.0'),
 ('产品使用文档与五场景教程',base+'Wiki/README.md'),
 ('系统架构与服务职责',base+'Wiki/architecture/overview.md'),
 ('规则编排与连续模拟',base+'Wiki/user-guide/definitions.md'),
 ('1.1 版本能力与验证范围',base+'Wiki/releases/1.1.md'),
 ('设备采集测量记录（2026-09-20）',base+'docs/evidence/capacity-disk-60m.json'),
 ('API 契约与外部集成',base+'Wiki/api/rest.md'),
 ('组件来源和许可材料',base+'deploy/licenses/README.md'),
]:link(label,url)
h('持续扩展方式',2)
p('新增设备可沿协议连接器与测点动作映射扩展，新的工艺场景可组合分析、告警和联动定义，外部系统通过公开 API、MCP 或进程插件协作。设备、规则、业务记录与运行版本共同保留来源，为工厂逐步增加场景和维护人员协作提供可追踪的基础。')

OUT.parent.mkdir(exist_ok=True)
for element in [doc._element,doc.styles._element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
